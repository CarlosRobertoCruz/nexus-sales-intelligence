"""Prepare local pilot references from original XLSX/ZIP reports, without DB/API access.

Requires Python + openpyxl. Does not modify workbooks or include customer PII in output.
"""
import argparse
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def key(value):
    if isinstance(value, bool):
        raise ValueError('CHAVE_INVALIDA')
    if isinstance(value, (float, int)):
        if not math.isfinite(value) or int(value) != value or value > 9007199254740991:
            raise ValueError('CHAVE_NUMERICA_SEM_PRECISAO')
        value = str(int(value))
    value = str(value)
    if not re.fullmatch(r'\d{1,40}', value):
        raise ValueError('CHAVE_INVALIDA')
    return value


def report_day(value):
    if value is None or value in ('', '-', '--'):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if re.match(r'^\d{2}/\d{2}/\d{4}(?:$|\s)', text):
        return datetime.strptime(text[:10], '%d/%m/%Y').date().isoformat()
    if re.match(r'^\d{4}-\d{2}-\d{2}(?:$|[ T])', text):
        return date.fromisoformat(text[:10]).isoformat()
    raise ValueError('DATA_RELATORIO_INVALIDA')


def read_report(path, identity, date_field):
    if path.suffix.lower() == '.zip':
        with zipfile.ZipFile(path) as archive:
            entries = [item for item in archive.infolist() if item.filename.lower().endswith('.xlsx')]
            if len(entries) != 1 or entries[0].file_size > 10 * 1024 * 1024:
                raise ValueError('ARQUIVO_RELATORIO_INVALIDO')
            filename = Path(entries[0].filename).name
            content = archive.read(entries[0])
    else:
        filename, content = path.name, path.read_bytes()
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('ARQUIVO_RELATORIO_INVALIDO')
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if len(book.worksheets) != 1:
            raise ValueError('ABA_RELATORIO_AMBIGUA')
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        if len(set(headers)) != len(headers) or identity not in headers or date_field not in headers:
            raise ValueError('COLUNAS_RELATORIO_INVALIDAS')
        records, seen, footers = [], set(), 0
        for excel_row, values in enumerate(rows, 2):
            if excel_row > 50002:
                raise ValueError('LIMITE_50000_REGISTROS')
            if all(value is None for value in values):
                continue
            record = dict(zip(headers, values))
            value = record[identity]
            # Hubsoft exports one final separator/footer row. Never skip an
            # arbitrary invalid identifier in the middle of the source table.
            if value == '-' and excel_row == sheet.max_row and report_day(record[date_field]) is None:
                footers += 1
                continue
            identifier = key(value)
            if identifier in seen:
                raise ValueError('CHAVES_RELATORIO_DUPLICADAS')
            seen.add(identifier)
            records.append({'key': identifier, 'day': report_day(record[date_field])})
        return records, {'filename': filename, 'fileSha256': hashlib.sha256(content).hexdigest(),
                         'sourceRecords': len(records), 'excludedFooters': footers}
    finally:
        book.close()


def prepare(args):
    pilot = date.fromisoformat(args.day)
    snapshot, followup = date.fromisoformat(args.snapshot_date), date.fromisoformat(args.followup_cutoff)
    if pilot.year != 2026 or pilot.month != 10 or not pilot <= followup < snapshot or followup.year != 2026:
        raise ValueError('PERIODO_PILOTO_INVALIDO')
    a, a_source = read_report(Path(args.attendance), 'protocolo', 'data_cadastro')
    o, o_source = read_report(Path(args.orders), 'numero_ordem_servico', 'data_termino_executado')
    references = {}
    for kind, records, source, field, identity, exported_field in [
        ('attendance', a, a_source, 'data_cadastro', 'protocolo', 'data_cadastro'),
        ('serviceOrders', o, o_source, 'data_termino_executado', 'numero_ordem_servico', args.os_export_date_field),
    ]:
        references[kind] = {**source, 'exportedAt': args.snapshot_date, 'exportedAtPrecision': 'day',
            'keyField': identity, 'keys': sorted(row['key'] for row in records if row['day'] == args.day),
            'filters': {'dateField': field, 'otherFilters': 'none'},
            'exportFilters': {'dateField': exported_field, 'start': '2026-10-01', 'end': '2026-10-31', 'otherFilters': 'none'}}
    references['attendance']['filters']['addressType'] = 'instalacao'
    references['attendance']['exportFilters']['addressType'] = 'instalacao'
    references['followup'] = {'pending': True, 'keyField': 'numero_ordem_servico',
        'filters': {'dateField': 'data_cadastro', 'otherFilters': 'none', 'relationship': 'id_atendimento', 'followupCutoff': args.followup_cutoff}}
    manifest = {'year': 2026, 'phase': 'day', 'peakDay': args.day, 'cutoff': args.day,
        'followupCutoff': args.followup_cutoff, 'osDateField': 'data_termino_executado', 'maxCalls': 500, 'reports': references}
    evidence = {'month': '2026-10', 'reviewer': None, 'confirmFilters': False, 'confirmIncompleteRecords': False,
        **{kind: {'reportSha256': references[kind]['fileSha256'], 'keys': references[kind]['keys']} for kind in ['attendance', 'serviceOrders']}}
    target = Path(args.output)
    target.mkdir(parents=True, exist_ok=True)
    for name, value in [('manifest.json', manifest), ('references.json', references), ('evidence-draft.json', evidence)]:
        path = target / name
        with path.open('x', encoding='utf-8') as output:
            json.dump(value, output, ensure_ascii=False, indent=2)
            output.write('\n')
    print(json.dumps({'pilot': args.day, 'attendanceReferenceCount': len(references['attendance']['keys']),
        'osReferenceCount': len(references['serviceOrders']['keys']), 'osExportDateField': args.os_export_date_field,
        'followupReference': 'pending', 'enabled': False}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--attendance', required=True)
    parser.add_argument('--orders', required=True)
    parser.add_argument('--day', required=True)
    parser.add_argument('--snapshot-date', required=True)
    parser.add_argument('--followup-cutoff', required=True)
    parser.add_argument('--os-export-date-field', required=True, choices=['data_cadastro', 'data_inicio_programado', 'data_termino_executado'])
    parser.add_argument('--output', required=True)
    prepare(parser.parse_args())
